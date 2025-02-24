from openpyxl import load_workbook

wb = load_workbook("sample/sample_v3.xlsx")

# シート一覧を表示
print("=== シート一覧 ===")
for sheet_name in wb.sheetnames:
    print(sheet_name)

# 各シートの最初の数行を確認
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== シート '{sheet_name}' の内容 ===")
    for row in range(1, 12):
        for cell in ws[row]:
            if cell.value:
                print(f"{cell.coordinate}: {cell.value}")
