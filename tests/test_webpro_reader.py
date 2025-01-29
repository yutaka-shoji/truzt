"""Test module for webpro reader."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from truzt.building_reader import BuildingReader
from truzt.reader import ExcelReadError
from truzt.room_reader import RoomReader
from truzt.webpro_reader import read_webpro_excel


def test_read_webpro_excel_v3(tmp_path: Path):
    """Test reading WEBPRO Excel file (v3)."""
    # サンプルファイル読み込み
    sample_file = "sample/sample_input_v3.xlsx"
    webpro = read_webpro_excel(sample_file, version="v3")

    # リファレンスJSONの読み込み
    with open("sample/sample_input_v3.json") as f:
        ref_data = json.load(f)

    # オブジェクトをdictに変換して比較
    webpro_dict = webpro.model_dump(by_alias=True)

    # 項目を個別に比較（順序に依存しない比較のため）
    assert json.dumps(
        webpro_dict["Building"], indent=2, ensure_ascii=False, sort_keys=True
    ) == json.dumps(ref_data["Building"], indent=2, ensure_ascii=False, sort_keys=True)
    assert json.dumps(
        webpro_dict["Rooms"], indent=2, ensure_ascii=False, sort_keys=True
    ) == json.dumps(ref_data["Rooms"], indent=2, ensure_ascii=False, sort_keys=True)

    # エラーケースの検証
    with pytest.raises(FileNotFoundError):
        read_webpro_excel("non_existent.xlsx")


def test_building_reader_v3():
    """Test BuildingReader."""
    # Excelファイルの読み込み
    wb = load_workbook("sample/sample_input_v3.xlsx", data_only=True)
    reader = BuildingReader(wb, version="v3")
    building = reader.read()

    # リファレンスJSONの読み込み
    with open("sample/sample_input_v3.json") as f:
        ref_data = json.load(f)
    ref_building_dict = ref_data["Building"]

    # オブジェクトをdictに変換して比較
    building_dict = building.model_dump(by_alias=True)
    assert json.dumps(building_dict, indent=2, ensure_ascii=False, sort_keys=True) == json.dumps(
        ref_building_dict, indent=2, ensure_ascii=False, sort_keys=True
    )


def test_room_reader_v3():
    """Test RoomReader."""
    # Excelファイルの読み込み
    wb = load_workbook("sample/sample_input_v3.xlsx", data_only=True)
    reader = RoomReader(wb, version="v3")
    rooms = reader.read()

    # リファレンスJSONの読み込み
    with open("sample/sample_input_v3.json") as f:
        ref_data = json.load(f)
    ref_rooms_dict = ref_data["Rooms"]

    # オブジェクトをdictに変換して比較
    rooms_dict = rooms.model_dump(by_alias=True)
    assert json.dumps(rooms_dict, indent=2, ensure_ascii=False, sort_keys=True) == json.dumps(
        ref_rooms_dict, indent=2, ensure_ascii=False, sort_keys=True
    )


def test_read_cell_error():
    """Test read_cell error handling."""
    wb = load_workbook("sample/sample_input_v3.xlsx", data_only=True)
    reader = BuildingReader(wb, version="v3")

    with pytest.raises(ExcelReadError):
        reader.read_cell("不存在シート", "A1")

    with pytest.raises(ExcelReadError):
        reader.read_cell("0) 基本情報", "Z999")  # 存在しないセル
