import json
from pathlib import Path

from openpyxl import load_workbook
from truzt.air_conditioning_zone_reader import AirConditioningZoneReader


def test_air_conditioning_zone_excel_read() -> None:
    """空調ゾーン情報のExcel読み込みテスト."""
    # テスト用のExcelファイルを読み込む
    wb = load_workbook(
        filename=Path("sample/sample_input_v3.xlsx"),
        data_only=True,
    )

    # AirConditioningZoneReaderを使用してゾーン情報を読み込む
    reader = AirConditioningZoneReader(wb, "v3")
    zones = reader.read()

    # 正解のJSONファイルを読み込む
    with open(Path("sample/sample_input_v3.json")) as file:
        data = json.load(file)

    # JSONから読み込んだデータと比較
    ac_zone_dict = data["AirConditioningZone"]
    ac_zone_loaded_dict = zones.model_dump(by_alias=True)

    assert ac_zone_dict == ac_zone_loaded_dict


def test_air_conditioning_zone_no_data_read() -> None:
    """空調ゾーン情報の読み込み（データなし）テスト."""
    # テスト用の空のExcelファイルを読み込む
    wb = load_workbook(
        filename=Path("src/truzt/data/blank_input_sheet_v3.6.xlsx"),
        data_only=True,
    )

    # AirConditioningZoneReaderを使用してゾーン情報を読み込む
    reader = AirConditioningZoneReader(wb, "v3")
    zones = reader.read()

    # 空の辞書になっていることを確認
    assert zones.model_dump(by_alias=True) == {}
